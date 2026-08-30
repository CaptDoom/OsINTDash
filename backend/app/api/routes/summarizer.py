import re
import hashlib
import logging
import httpx
import pypdf
import docx
from datetime import datetime, timedelta, timezone
from io import BytesIO
from typing import List, Optional
from fastapi import APIRouter, Depends, HTTPException, UploadFile, File, Form

logger = logging.getLogger("drishya.api.summarizer")
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession
from pydantic import BaseModel

from backend.app.database import get_db, Article
from backend.app.config import settings
from backend.app.services.summarizer import call_openai, call_gemini, call_ollama
from backend.app.redis_pool import cache_get, cache_set

router = APIRouter(prefix="/api/summarizer")

COUNTRY_NAMES_BY_CODE = {
    "CN": "China",
    "PK": "Pakistan",
    "AF": "Afghanistan",
    "BD": "Bangladesh",
    "MM": "Myanmar",
    "NP": "Nepal",
    "BT": "Bhutan",
    "LK": "Sri Lanka",
    "MV": "Maldives",
    "IN": "India",
    "US": "United States",
    "RU": "Russia",
    "IR": "Iran",
    "IL": "Israel",
    "TW": "Taiwan",
    "JP": "Japan",
    "UA": "Ukraine",
}

def parse_pdf(file_bytes: bytes) -> str:
    try:
        reader = pypdf.PdfReader(BytesIO(file_bytes))
        text = ""
        for page in reader.pages:
            text += (page.extract_text() or "") + "\n"
        return text
    except Exception as e:
        return f"[Error parsing PDF: {str(e)}]"

def parse_docx(file_bytes: bytes) -> str:
    try:
        doc = docx.Document(BytesIO(file_bytes))
        return "\n".join([p.text for p in doc.paragraphs])
    except Exception as e:
        return f"[Error parsing Word Document: {str(e)}]"

async def scrape_url(url: str) -> str:
    from backend.app.services.net_safety import is_safe_url

    # SSRF protection: reject URLs pointing to private/loopback/link-local networks
    if not is_safe_url(url):
        logger.warning("[Summarizer] SSRF blocked: %s", url)
        return f"URL Source ({url}) blocked: local/private network access is restricted."

    # Check Redis cache for previously scraped URLs (1-hour TTL)
    url_hash = hashlib.md5(url.encode()).hexdigest()
    cache_key = f"drishya:scrape:url:{url_hash}"
    cached = await cache_get(cache_key)
    if cached and isinstance(cached, str):
        return cached
    try:
        # follow_redirects=False + manual hop validation to prevent SSRF-via-redirect
        async with httpx.AsyncClient(follow_redirects=False) as client:
            headers = {"User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36"}
            resp = await client.get(url, headers=headers, timeout=10.0)
            hops = 0
            while resp.is_redirect and hops < 3:
                next_url = str(resp.next_request.url)
                if not is_safe_url(next_url):
                    return f"URL Source ({url}) blocked: redirect target is a restricted address."
                resp = await client.get(next_url, headers=headers)
                hops += 1
            if resp.status_code == 200:
                html = resp.text
                html = re.sub(r'<script.*?>.*?</script>', '', html, flags=re.DOTALL | re.IGNORECASE)
                html = re.sub(r'<style.*?>.*?</style>', '', html, flags=re.DOTALL | re.IGNORECASE)
                text = re.sub(r'<[^>]+>', ' ', html)
                text = re.sub(r'\s+', ' ', text).strip()
                result = f"URL Source ({url}):\n{text[:3000]}"
                # Cache for 1 hour
                await cache_set(cache_key, result, ttl=3600)
                return result
            return f"URL Source ({url}) failed with status {resp.status_code}"
    except Exception as e:
        return f"URL Source ({url}) failed: {str(e)}"

def generate_fallback_summary(country_name: str, timeframe_label: str, articles: list, external_context: str) -> str:
    md = f"# News Briefing: {country_name} ({timeframe_label})\n\n"
    
    md += "## Summary\n"
    if articles:
        sources = set(art.source for art in articles if art.source)
        sources_str = ", ".join(list(sources)[:3])
        headlines_summary = "; ".join(f"'{art.title}'" for art in articles[:2])
        md += f"{len(articles)} updates found for {country_name} over the last {timeframe_label} from **{sources_str or 'news sources'}**. Key reports: {headlines_summary}. Activity is within normal levels.\n\n"
    else:
        md += f"No significant updates for {country_name} over the last {timeframe_label}. Things are quiet and stable.\n\n"
    
    md += "## Details by Area\n"
    
    # Group by dept
    by_dept = {}
    for art in articles:
        dept = art.department
        if dept not in by_dept:
            by_dept[dept] = []
        by_dept[dept].append(art)
        
    for dept in ["Military & Defense", "Economic & Financial", "Political & Diplomatic", "Social Affairs & Welfare / Technology"]:
        md += f"### {dept}\n"
        # Match departments
        if dept == "Social Affairs & Welfare / Technology":
            dept_arts = by_dept.get("Social Affairs & Welfare", []) + by_dept.get("Technology & Cyber", [])
        else:
            dept_arts = by_dept.get(dept, [])
            
        if not dept_arts:
            md += f"- No reports in this area. Things are normal.\n\n"
        else:
            src_list = list(set(art.source for art in dept_arts if art.source))[:3]
            srcs_str = " and ".join(src_list) if len(src_list) > 1 else (src_list[0] if src_list else "news sources")
            md += f"- Recent updates from **{srcs_str}**:\n"
            for art in dept_arts[:4]:
                md += f"  - **[{art.title}]({art.url})** ({art.source}): {art.summary or art.content[:180]}\n"
            md += "\n"
            
    if external_context.strip() and "No external documents" not in external_context:
        md += "### Additional Details\n"
        md += f"> {external_context[:800]}\n\n"
        
    md += "## What This Means for People\n"
    if articles:
        md += f"Daily life, travel, and safety are not affected. Everything continues as normal.\n"
    else:
        md += "Nothing unusual. Life goes on as normal. Check back for updates.\n"
    return md

@router.post("/generate")
async def generate_custom_summary(
    country_code: str = Form(...),
    timeframe: str = Form(...),
    urls: Optional[str] = Form(None),
    files: Optional[List[UploadFile]] = File(None),
    db: AsyncSession = Depends(get_db)
):
    """
    Generate an AI-driven geopolitical summary of a country over a timeframe (1M, 6M, 1Y)
    blended with external context from uploaded documents and links.
    """
    country_name = COUNTRY_NAMES_BY_CODE.get(country_code.upper(), country_code)
    
    # Calculate cutoff time based on timeframe
    now = datetime.now(timezone.utc)
    if timeframe == "1M":
        start_date = now - timedelta(days=30)
        timeframe_label = "1 Month"
    elif timeframe == "6M":
        start_date = now - timedelta(days=180)
        timeframe_label = "6 Months"
    elif timeframe == "1Y":
        start_date = now - timedelta(days=365)
        timeframe_label = "1 Year"
    else:
        raise HTTPException(status_code=400, detail="Invalid timeframe. Must be '1M', '6M', or '1Y'.")

    # 1. Fetch articles from database
    stmt = select(Article).where(
        Article.country_code == country_code.upper(),
        Article.published_at >= start_date
    ).order_by(Article.published_at.desc())
    
    result = await db.execute(stmt)
    articles = list(result.scalars().all())

    # Formulate articles context
    db_articles_context = ""
    if articles:
        for idx, art in enumerate(articles[:40]): # Limit to 40 articles to prevent token overflow
            db_articles_context += (
                f"SOURCE [{idx+1}]: {art.title}\n"
                f"URL: {art.url}\n"
                f"Department: {art.department}\n"
                f"Source: {art.source}\n"
                f"Published: {art.published_at.isoformat()}\n"
                f"Summary: {art.summary or art.content[:200]}\n\n"
            )
    else:
        db_articles_context = "No news reports found in database for this timeframe.\n"

    # 2. Parse external context
    external_texts = []
    
    # Scrape web URLs
    if urls:
        url_list = [u.strip() for u in urls.split(",") if u.strip()]
        for url in url_list:
            scraped = await scrape_url(url)
            external_texts.append(scraped)

    # Parse uploaded files
    if files:
        for upload_file in files:
            file_bytes = await upload_file.read()
            filename = upload_file.filename.lower()
            if filename.endswith(".pdf"):
                parsed = parse_pdf(file_bytes)
            elif filename.endswith((".docx", ".doc")):
                parsed = parse_docx(file_bytes)
            else:
                # Text files
                try:
                    parsed = file_bytes.decode("utf-8", errors="ignore")
                except Exception:
                    parsed = f"[Error decoding text file {upload_file.filename}]"
            
            external_texts.append(f"Document ({upload_file.filename}):\n{parsed}")

    external_context = "\n\n---\n\n".join(external_texts) if external_texts else "No external documents or web links provided."

    # 3. LLM Generation
    prompt = f"""
ROLE:
You are an expert communicator who translates complex news into simple, plain English. Your goal is to explain what happened in clear terms without using any technical, political, or military jargon.

GOAL:
Eliminate superficial summaries, empty placeholders, and technical terminology. Every summary must be clear, easy to understand, and richly structured in plain language.

---

NEWS AND STABILITY BRIEFING: {country_name.upper()}
TIMEFRAME: {timeframe_label}

INPUT SOURCES:
1. INTERNAL NEWS DATABASE:
-----------------------------------------------------
{db_articles_context}
-----------------------------------------------------

2. USER-UPLOADED DOCUMENTS & WEB LINKS (External context):
-----------------------------------------------------
{external_context}
-----------------------------------------------------

---

### INSTRUCTION SET & STANDARDS

1. SIMPLIFIED LANGUAGE:
   - Use simple, everyday words. Avoid any jargon, such as "OSINT," "telemetry," "bilaterals," "strategic meetings," "tactical," "reconnaissance," "frontier," etc.
   - Explain everything in plain English so an average person can easily understand it.
   - Do not use conversational filler (e.g., "In conclusion," "It is important to note").

2. MANDATORY STRUCTURE:
   Unless specified otherwise, every comprehensive summary must include:
   - **1. SUMMARY**: A 2-3 sentence overview explaining what is happening and how it affects general stability.
   - **2. SECTOR DETAILS**: Group key information under distinct, thematic headings:
     - **Military & Defense** (explain guard work, safety drills, or patrol adjustments in simple terms)
     - **Economic & Financial** (explain trade, prices, or infrastructure developments in simple terms)
     - **Political & Diplomatic** (explain leadership meetings or agreements in simple terms)
     - **Social Affairs & Welfare / Technology** (explain community welfare, clinics, or network safety in simple terms)
   - **3. WHAT THIS MEANS FOR ORDINARY PEOPLE**: Explain in 1-2 sentences how this directly affects average citizen safety, costs, travel, or stability.

3. ACCURACY FALLBACK PROTOCOL:
   - If direct data for a sector is sparse in the provided context, describe the general situation using simple terms.
   - If information is missing, highlight it as a "Missing Info" note.

4. FORMATTING RULES:
   - Use scannable markdown: Bold key entities, utilize simple bullet points for lists.
   - Maintain a direct, objective, and clear tone.
   - Highlight links and citations from the database articles where appropriate.

---

### INPUT CONTEXT PROTOCOL
1. Extract primary entities, key metrics, and time-bound events.
2. Cross-reference provided external sources (PDFs, URLs, Notes) with internal context to create a unified narrative.
"""

    summary_text = ""
    # Try calling available LLMs
    if settings.llm_provider == "ollama" and settings.ollama_base_url:
        try:
            summary_text = await call_ollama(prompt, "You are a Senior Intel Fusion Officer.")
        except Exception:
            pass

    if not summary_text:
        if settings.openai_api_key:
            try:
                summary_text = await call_openai(prompt, "You are a Senior Intel Fusion Officer.")
            except Exception:
                pass
        elif settings.google_api_key:
            try:
                summary_text = await call_gemini(prompt, "You are a Senior Intel Fusion Officer.")
            except Exception:
                pass
                
    if not summary_text:
        # Heuristic fallback summary if all LLMs are offline or not configured
        summary_text = generate_fallback_summary(country_name, timeframe_label, articles, external_context)

    return {"summary": summary_text}


@router.post("/stream")
async def generate_streaming_summary(
    country_code: str = Form(...),
    timeframe: str = Form(...),
    urls: Optional[str] = Form(None),
    db: AsyncSession = Depends(get_db)
):
    """
    Streaming summarizer using Server-Sent Events.
    Returns article metadata immediately, then streams the LLM briefing in real-time.
    """
    from fastapi.responses import StreamingResponse
    import json as _json

    country_name = COUNTRY_NAMES_BY_CODE.get(country_code.upper(), country_code)

    # Calculate cutoff
    now = datetime.now(timezone.utc)
    if timeframe == "1M":
        start_date, timeframe_label = now - timedelta(days=30), "1 Month"
    elif timeframe == "6M":
        start_date, timeframe_label = now - timedelta(days=180), "6 Months"
    elif timeframe == "1Y":
        start_date, timeframe_label = now - timedelta(days=365), "1 Year"
    else:
        raise HTTPException(status_code=400, detail="Invalid timeframe. Must be '1M', '6M', or '1Y'.")

    # Fetch articles
    stmt = select(Article).where(
        Article.country_code == country_code.upper(),
        Article.published_at >= start_date
    ).order_by(Article.published_at.desc())
    result = await db.execute(stmt)
    articles = list(result.scalars().all())

    # Build article metadata for frontend
    article_meta = []
    db_articles_context = ""
    for idx, art in enumerate(articles[:40]):
        db_articles_context += (
            f"SOURCE [{idx+1}]: {art.title}\n"
            f"URL: {art.url}\n"
            f"Department: {art.department}\n"
            f"Source: {art.source}\n"
            f"Impact: {art.impact_level}\n"
            f"Published: {art.published_at.isoformat()}\n"
            f"Summary: {art.summary or art.content[:200]}\n\n"
        )
        article_meta.append({
            "id": art.id, "title": art.title, "url": art.url,
            "source": art.source, "department": art.department,
            "impact_level": art.impact_level,
            "published_at": art.published_at.isoformat(),
        })

    # Parse external URLs if provided
    external_context = "No external sources provided."
    if urls:
        url_list = [u.strip() for u in urls.split(",") if u.strip()]
        ext_texts = []
        for url in url_list:
            scraped = await scrape_url(url)
            ext_texts.append(scraped)
        if ext_texts:
            external_context = "\n\n---\n\n".join(ext_texts)

    # Build prompt
    prompt = f"""
You are an expert communicator who translates complex news into simple, plain English.
Write a cohesive news and stability briefing for {country_name.upper()} over {timeframe_label}.

NEWS REPORTS:
{db_articles_context}

EXTERNAL SOURCES:
{external_context}

FORMAT:
**1. Summary** - 2-3 sentence overview.
**2. Sector Details** - Group by: Military & Defense, Economic & Financial, Political & Diplomatic, Social & Welfare / Technology.
**3. What This Means for Ordinary People** - Plain English impact.
**4. Source Reliability** - Note which sources reported each key finding.

RULES:
- No jargon (no OSINT, telemetry, bilaterals, tactical, reconnaissance, frontier, strategic, etc.).
- Be direct. Every sentence must add new information.
- Keep under 400 words.
"""

    async def event_generator():
        # First event: article metadata + stats
        high = sum(1 for a in articles if a.impact_level == "High Impact")
        medium = sum(1 for a in articles if a.impact_level == "Medium Impact")
        sources = list(set(a.source for a in articles if a.source))[:10]
        meta_data = {
            'country': country_name,
            'timeframe': timeframe_label,
            'total_articles': len(articles),
            'high_impact': high,
            'medium_impact': medium,
            'sources': sources,
            'articles': article_meta[:20],
        }
        yield f"event: metadata\ndata: {_json.dumps(meta_data)}\n\n"

        # Stream LLM response
        try:
            full_text = ""
            if settings.openai_api_key:
                from openai import AsyncOpenAI
                client = AsyncOpenAI(api_key=settings.openai_api_key)
                stream = await client.chat.completions.create(
                    model="gpt-4o-mini",
                    messages=[
                        {"role": "system", "content": "You are a Senior Intel Fusion Officer providing clear, plain-English briefings."},
                        {"role": "user", "content": prompt},
                    ],
                    temperature=0.2,
                    stream=True,
                )
                async for chunk in stream:
                    delta = chunk.choices[0].delta.content if chunk.choices else None
                    if delta:
                        full_text += delta
                        yield f"event: token\ndata: {_json.dumps({'text': delta})}\n\n"
            else:
                if settings.llm_provider == "ollama" and settings.ollama_base_url:
                    try:
                        full_text = await call_ollama(prompt, "You are a clear and simple writer.")
                    except Exception:
                        pass
                if not full_text and settings.google_api_key:
                    try:
                        full_text = await call_gemini(prompt, "You are a clear and simple writer.")
                    except Exception:
                        pass
                if not full_text:
                    full_text = generate_fallback_summary(country_name, timeframe_label, articles, external_context)
                yield f"event: token\ndata: {_json.dumps({'text': full_text})}\n\n"

        except Exception as exc:
            logger.error("[Summarizer Stream] LLM call failed: %s", exc)
            full_text = generate_fallback_summary(country_name, timeframe_label, articles, external_context)
            yield f"event: token\ndata: {_json.dumps({'text': full_text})}\n\n"

        yield f"event: done\ndata: {_json.dumps({'summary': full_text})}\n\n"

    return StreamingResponse(event_generator(), media_type="text/event-stream")
